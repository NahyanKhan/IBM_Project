"""
convert_report.py
Converts MohammedNahyanKhan_ProjectReport.md → MohammedNahyanKhan_ProjectReport.docx
using python-docx only (no Pandoc binary required).
Run: python convert_report.py
"""

import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MD_FILE   = "MohammedNahyanKhan_ProjectReport.md"
DOCX_FILE = "MohammedNahyanKhan_ProjectReport.docx"

# ── Helpers ────────────────────────────────────────────────────────────────────

def set_run_formatting(run, text):
    """Apply bold/italic/code inline formatting from markdown spans."""
    # inline code  `text`
    if text.startswith("`") and text.endswith("`"):
        run.text = text[1:-1]
        run.font.name = "Courier New"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
        return
    # bold+italic ***text***
    if text.startswith("***") and text.endswith("***"):
        run.text = text[3:-3]
        run.bold = True
        run.italic = True
        return
    # bold **text**
    if text.startswith("**") and text.endswith("**"):
        run.text = text[2:-2]
        run.bold = True
        return
    # italic *text*
    if text.startswith("*") and text.endswith("*"):
        run.text = text[1:-1]
        run.italic = True
        return
    run.text = text


def add_formatted_paragraph(doc, text, style="Normal", alignment=None):
    """
    Add a paragraph that supports inline **bold**, *italic*, `code` spans.
    Returns the paragraph object.
    """
    para = doc.add_paragraph(style=style)
    if alignment is not None:
        para.alignment = alignment

    # Split on bold/italic/code markers while keeping delimiters
    tokens = re.split(r"(\*\*\*[^*]+\*\*\*|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text)
    for token in tokens:
        if not token:
            continue
        run = para.add_run()
        set_run_formatting(run, token)
    return para


def add_code_block(doc, text):
    """Add a monospaced code block paragraph."""
    para = doc.add_paragraph(style="Normal")
    run = para.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(9)
    para.paragraph_format.left_indent = Inches(0.4)
    # light grey shading
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F2F2F2")
    pPr.append(shd)
    return para


def add_horizontal_rule(doc):
    """Simulate an --- divider with a bottom border on an empty paragraph."""
    para = doc.add_paragraph()
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "CCCCCC")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return para


# ── Main conversion ────────────────────────────────────────────────────────────

doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin    = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin   = Inches(1.2)
    section.right_margin  = Inches(1.2)

with open(MD_FILE, encoding="utf-8") as f:
    lines = f.readlines()

in_code_block = False
code_lines    = []

for raw_line in lines:
    line = raw_line.rstrip("\n")

    # ── Code block fence ──────────────────────────────────────────────────────
    if line.strip().startswith("```"):
        if not in_code_block:
            in_code_block = True
            code_lines = []
        else:
            in_code_block = False
            add_code_block(doc, "\n".join(code_lines))
        continue

    if in_code_block:
        code_lines.append(line)
        continue

    # ── Horizontal rule ───────────────────────────────────────────────────────
    if re.match(r"^-{3,}$", line.strip()):
        add_horizontal_rule(doc)
        continue

    # ── Headings ──────────────────────────────────────────────────────────────
    h1 = re.match(r"^# (.+)$", line)
    h2 = re.match(r"^## (.+)$", line)
    h3 = re.match(r"^### (.+)$", line)

    if h1:
        p = add_formatted_paragraph(doc, h1.group(1), style="Heading 1")
        continue
    if h2:
        add_formatted_paragraph(doc, h2.group(1), style="Heading 2")
        continue
    if h3:
        add_formatted_paragraph(doc, h3.group(1), style="Heading 3")
        continue

    # ── Screenshot placeholder (blockquote) ───────────────────────────────────
    bq = re.match(r"^>\s*\*?\*?\[(.+)\]\*?\*?$", line.strip())
    if bq:
        para = doc.add_paragraph(style="Normal")
        para.paragraph_format.left_indent = Inches(0.4)
        run = para.add_run(f"[ SCREENSHOT PLACEHOLDER: {bq.group(1)} ]")
        run.italic = True
        run.font.color.rgb = RGBColor(0x99, 0x66, 0x00)
        run.font.size = Pt(10)
        continue

    # ── Bullet points ─────────────────────────────────────────────────────────
    bullet = re.match(r"^- (.+)$", line)
    if bullet:
        add_formatted_paragraph(doc, bullet.group(1), style="List Bullet")
        continue

    # ── Metadata lines (bold key: value at top of doc) ────────────────────────
    meta = re.match(r"^\*\*(.+?):\*\* (.+)$", line)
    if meta:
        para = doc.add_paragraph(style="Normal")
        para.add_run(meta.group(1) + ": ").bold = True
        run_bold = para.runs[-1]
        run_bold.bold = True
        para.add_run(meta.group(2))
        continue

    # ── Empty line → spacing ──────────────────────────────────────────────────
    if line.strip() == "":
        doc.add_paragraph()
        continue

    # ── Default: normal paragraph with inline formatting ─────────────────────
    # Strip leading blockquote marker if any
    clean = re.sub(r"^>\s*", "", line)
    add_formatted_paragraph(doc, clean)

doc.save(DOCX_FILE)
print(f"Created: {DOCX_FILE}")
